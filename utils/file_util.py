#!/usr/bin/env python3
"""
-------------------------------------------------
File:           file_util.py
Author:         duanyang
Date:           2025/11/27
-------------------------------------------------
Description:    
文件处理模块,提供常用的文件操作函数,如复制、移动、重命名、删除文件等
-------------------------------------------------
"""
import csv
import hashlib
import json
import os
import shutil
import time

import xlrd
import yaml
from openpyxl import load_workbook

from config.config_manager import config_manager
from utils.logger_util import logger


class FileHandler:
    """
    通用文件和数据处理类,用于文件和数据的全面操作和管理。
    """

    @staticmethod
    def read_json(file_path):
        """
        读取JSON文件
        
        Args:
            file_path: JSON文件路径,可以是相对路径或绝对路径
            
        Returns:
            dict/list: 解析后的JSON数据
            
        Raises:
            FileNotFoundError: 文件不存在
            json.JSONDecodeError: JSON格式错误
        """
        # 处理相对路径
        if not os.path.isabs(file_path):
            file_path = os.path.join(config_manager.DATA_DIR, file_path)

        logger.debug(f"读取JSON文件: {file_path}")
        try:
            with open(file_path, 'r', encoding = 'utf-8') as f:
                return json.load(f)
        except FileNotFoundError:
            logger.error(f"JSON文件不存在: {file_path}")
            raise
        except json.JSONDecodeError as e:
            logger.error(f"JSON格式错误: {file_path}, 错误: {str(e)}")
            raise

    @staticmethod
    def write_json(file_path, data, indent = 2):
        """
        写入JSON文件
        
        Args:
            file_path: JSON文件路径,可以是相对路径或绝对路径
            data: 要写入的数据
            indent: 缩进空格数
        """
        # 处理相对路径
        if not os.path.isabs(file_path):
            file_path = os.path.join(config_manager.DATA_DIR, file_path)

        # 确保目录存在
        os.makedirs(os.path.dirname(file_path), exist_ok = True)

        logger.debug(f"写入JSON文件: {file_path}")
        try:
            with open(file_path, 'w', encoding = 'utf-8') as f:
                json.dump(data, f, ensure_ascii = False, indent = indent)
        except Exception as e:
            logger.error(f"写入JSON文件失败: {file_path}, 错误: {str(e)}")
            raise

    @staticmethod
    def read_yaml(file_path):
        """
        读取YAML文件
        
        Args:
            file_path: YAML文件路径,可以是相对路径或绝对路径
            
        Returns:
            dict/list: 解析后的YAML数据
        """
        # 处理相对路径
        if not os.path.isabs(file_path):
            file_path = os.path.join(config_manager.DATA_DIR, file_path)

        logger.debug(f"读取YAML文件: {file_path}")
        try:
            with open(file_path, 'r', encoding = 'utf-8') as f:
                return yaml.safe_load(f)
        except FileNotFoundError:
            logger.error(f"YAML文件不存在: {file_path}")
            raise
        except yaml.YAMLError as e:
            logger.error(f"YAML格式错误: {file_path}, 错误: {str(e)}")
            raise

    @staticmethod
    def write_yaml(file_path, data):
        """
        写入YAML文件
        
        Args:
            file_path: YAML文件路径,可以是相对路径或绝对路径
            data: 要写入的数据
        """
        # 处理相对路径
        if not os.path.isabs(file_path):
            file_path = os.path.join(config_manager.DATA_DIR, file_path)

        # 确保目录存在
        os.makedirs(os.path.dirname(file_path), exist_ok = True)

        logger.debug(f"写入YAML文件: {file_path}")
        try:
            with open(file_path, 'w', encoding = 'utf-8') as f:
                yaml.dump(data, f, default_flow_style = False, allow_unicode = True)
        except Exception as e:
            logger.error(f"写入YAML文件失败: {file_path}, 错误: {str(e)}")
            raise

    @staticmethod
    def read_csv(file_path):
        """
        读取CSV文件
        
        Args:
            file_path: CSV文件路径,可以是相对路径或绝对路径
            
        Returns:
            list: CSV数据列表
        """
        # 处理相对路径
        if not os.path.isabs(file_path):
            file_path = os.path.join(config_manager.DATA_DIR, file_path)

        logger.debug(f"读取CSV文件: {file_path}")
        try:
            data = []
            with open(file_path, 'r', encoding = 'utf-8') as f:
                reader = csv.DictReader(f)
                for row in reader:
                    data.append(dict(row))
            return data
        except FileNotFoundError:
            logger.error(f"CSV文件不存在: {file_path}")
            raise
        except Exception as e:
            logger.error(f"读取CSV文件失败: {file_path}, 错误: {str(e)}")
            raise

    @staticmethod
    def read_excel(file_path, sheet_name = None):
        """
        读取Excel文件
        
        Args:
            file_path: Excel文件路径,可以是相对路径或绝对路径
            sheet_name: 工作表名称,默认为第一个工作表
            
        Returns:
            list: Excel数据列表
        """
        # 处理相对路径
        if not os.path.isabs(file_path):
            file_path = os.path.join(config_manager.DATA_DIR, file_path)

        logger.debug(f"读取Excel文件: {file_path}")
        try:
            # 使用openpyxl读取xlsx格式
            if file_path.endswith('.xlsx'):
                workbook = load_workbook(file_path)
                if sheet_name:
                    sheet = workbook[sheet_name]
                else:
                    sheet = workbook.active

                # 获取表头
                headers = [cell.value for cell in sheet[1]]
                data = []

                # 读取数据行
                for row in sheet.iter_rows(min_row = 2, values_only = True):
                    data.append(dict(zip(headers, row)))
                return data
            # 使用xlrd读取xls格式
            elif file_path.endswith('.xls'):
                workbook = xlrd.open_workbook(file_path)
                if sheet_name:
                    sheet = workbook.sheet_by_name(sheet_name)
                else:
                    sheet = workbook.sheet_by_index(0)

                headers = [sheet.cell_value(0, col) for col in range(sheet.ncols)]
                data = []

                for row_idx in range(1, sheet.nrows):
                    row_data = {}
                    for col_idx in range(sheet.ncols):
                        row_data[headers[col_idx]] = sheet.cell_value(row_idx, col_idx)
                    data.append(row_data)
                return data
            else:
                raise ValueError(f"不支持的Excel格式: {file_path}")
        except FileNotFoundError:
            logger.error(f"Excel文件不存在: {file_path}")
            raise
        except Exception as e:
            logger.error(f"读取Excel文件失败: {file_path}, 错误: {str(e)}")
            raise

    @staticmethod
    def get_test_data(data_file, data_id = None):
        """
        获取测试数据
        根据文件扩展名自动选择读取方法
        
        Args:
            data_file: 数据文件路径
            data_id: 数据ID,用于从数据文件中筛选特定数据
            
        Returns:
            dict/list: 测试数据
        """
        # 根据文件扩展名选择读取方法
        ext = os.path.splitext(data_file)[1].lower()

        if ext == '.json':
            data = FileHandler.read_json(data_file)
        elif ext in ('.yaml', '.yml'):
            data = FileHandler.read_yaml(data_file)
        elif ext == '.csv':
            data = FileHandler.read_csv(data_file)
        elif ext in ('.xlsx', '.xls'):
            data = FileHandler.read_excel(data_file)
        else:
            raise ValueError(f"不支持的数据文件格式: {ext}")

        # 如果指定了data_id,则返回对应的数据
        if data_id is not None:
            if isinstance(data, list):
                for item in data:
                    if item.get('id') == data_id or item.get('case_id') == data_id:
                        return item
                raise ValueError(f"未找到ID为 {data_id} 的数据")
            elif isinstance(data, dict):
                if data_id in data:
                    return data[data_id]
                else:
                    raise ValueError(f"未找到ID为 {data_id} 的数据")

        return data

    @staticmethod
    def copy_file(src_path, dst_path):
        """
        复制文件
        
        Args:
            src_path: 源文件路径
            dst_path: 目标文件路径
        """
        # 直接使用已导入的config_manager
        # 处理相对路径
        if not os.path.isabs(src_path):
            src_path = os.path.join(config_manager.DATA_DIR, src_path)
        if not os.path.isabs(dst_path):
            dst_path = os.path.join(config_manager.DATA_DIR, dst_path)

        logger.debug(f"复制文件: {src_path} -> {dst_path}")
        try:
            # 确保目标目录存在
            os.makedirs(os.path.dirname(dst_path), exist_ok = True)
            shutil.copy2(src_path, dst_path)
            logger.info(f"文件复制成功: {src_path} -> {dst_path}")
        except FileNotFoundError:
            logger.error(f"源文件不存在: {src_path}")
            raise
        except Exception as e:
            logger.error(f"复制文件失败: {src_path} -> {dst_path}, 错误: {str(e)}")
            raise

    @staticmethod
    def move_file(src_path, dst_path):
        """
        移动文件
        
        Args:
            src_path: 源文件路径
            dst_path: 目标文件路径
        """
        # 直接使用已导入的config_manager
        # 处理相对路径
        if not os.path.isabs(src_path):
            src_path = os.path.join(config_manager.DATA_DIR, src_path)
        if not os.path.isabs(dst_path):
            dst_path = os.path.join(config_manager.DATA_DIR, dst_path)

        logger.debug(f"移动文件: {src_path} -> {dst_path}")
        try:
            # 确保目标目录存在
            os.makedirs(os.path.dirname(dst_path), exist_ok = True)
            shutil.move(src_path, dst_path)
            logger.info(f"文件移动成功: {src_path} -> {dst_path}")
        except FileNotFoundError:
            logger.error(f"源文件不存在: {src_path}")
            raise
        except Exception as e:
            logger.error(f"移动文件失败: {src_path} -> {dst_path}, 错误: {str(e)}")
            raise

    @staticmethod
    def rename_file(file_path, new_name):
        """
        重命名文件
        
        Args:
            file_path: 原文件路径
            new_name: 新文件名（不含路径）
        """
        # 直接使用已导入的config_manager
        if not os.path.isabs(file_path):
            file_path = os.path.join(config_manager.DATA_DIR, file_path)

        # 获取目录路径
        dir_path = os.path.dirname(file_path)
        new_path = os.path.join(dir_path, new_name)

        logger.debug(f"重命名文件: {file_path} -> {new_path}")
        try:
            os.rename(file_path, new_path)
            logger.info(f"文件重命名成功: {file_path} -> {new_path}")
        except FileNotFoundError:
            logger.error(f"文件不存在: {file_path}")
            raise
        except Exception as e:
            logger.error(f"重命名文件失败: {file_path} -> {new_path}, 错误: {str(e)}")
            raise

    @staticmethod
    def delete_file(file_path):
        """
        删除文件
        
        Args:
            file_path: 文件路径
        """
        # 直接使用已导入的config_manager
        # 处理相对路径
        if not os.path.isabs(file_path):
            file_path = os.path.join(config_manager.DATA_DIR, file_path)

        logger.debug(f"删除文件: {file_path}")
        try:
            os.remove(file_path)
            logger.info(f"文件删除成功: {file_path}")
        except FileNotFoundError:
            logger.error(f"文件不存在: {file_path}")
            raise
        except Exception as e:
            logger.error(f"删除文件失败: {file_path}, 错误: {str(e)}")
            raise

    @staticmethod
    def create_directory(dir_path):
        """
        创建目录
        
        Args:
            dir_path: 目录路径
        """
        # 直接使用已导入的config_manager
        # 处理相对路径
        if not os.path.isabs(dir_path):
            dir_path = os.path.join(config_manager.DATA_DIR, dir_path)

        logger.debug(f"创建目录: {dir_path}")
        try:
            os.makedirs(dir_path, exist_ok = True)
            logger.info(f"目录创建成功: {dir_path}")
        except Exception as e:
            logger.error(f"创建目录失败: {dir_path}, 错误: {str(e)}")
            raise

    @staticmethod
    def list_directory(dir_path):
        """
        列出目录内容
        
        Args:
            dir_path: 目录路径
            
        Returns:
            dict: 包含文件列表和目录列表
        """
        # 直接使用已导入的config_manager
        # 处理相对路径
        if not os.path.isabs(dir_path):
            dir_path = os.path.join(config_manager.DATA_DIR, dir_path)

        logger.debug(f"列出目录内容: {dir_path}")
        try:
            files = []
            directories = []

            for item in os.listdir(dir_path):
                item_path = os.path.join(dir_path, item)
                if os.path.isfile(item_path):
                    files.append(item)
                elif os.path.isdir(item_path):
                    directories.append(item)

            return {
                'files': files,
                'directories': directories
            }
        except FileNotFoundError:
            logger.error(f"目录不存在: {dir_path}")
            raise
        except Exception as e:
            logger.error(f"列出目录内容失败: {dir_path}, 错误: {str(e)}")
            raise

    @staticmethod
    def delete_directory(dir_path, recursive = False):
        """
        删除目录
        
        Args:
            dir_path: 目录路径
            recursive: 是否递归删除,True表示删除目录及其所有内容
        """
        # 直接使用已导入的config_manager
        # 处理相对路径
        if not os.path.isabs(dir_path):
            dir_path = os.path.join(config_manager.DATA_DIR, dir_path)

        logger.debug(f"删除目录: {dir_path}, 递归: {recursive}")
        try:
            if recursive:
                shutil.rmtree(dir_path)
                logger.info(f"递归删除目录成功: {dir_path}")
            else:
                os.rmdir(dir_path)
                logger.info(f"删除空目录成功: {dir_path}")
        except FileNotFoundError:
            logger.error(f"目录不存在: {dir_path}")
            raise
        except OSError:
            logger.error(f"目录不为空: {dir_path}, 使用recursive=True参数递归删除")
            raise
        except Exception as e:
            logger.error(f"删除目录失败: {dir_path}, 错误: {str(e)}")
            raise

    @staticmethod
    def copy_directory(src_path, dst_path):
        """
        复制目录
        
        Args:
            src_path: 源目录路径
            dst_path: 目标目录路径
        """
        # 直接使用已导入的config_manager
        # 处理相对路径
        if not os.path.isabs(src_path):
            src_path = os.path.join(config_manager.DATA_DIR, src_path)
        if not os.path.isabs(dst_path):
            dst_path = os.path.join(config_manager.DATA_DIR, dst_path)

        logger.debug(f"复制目录: {src_path} -> {dst_path}")
        try:
            shutil.copytree(src_path, dst_path)
            logger.info(f"目录复制成功: {src_path} -> {dst_path}")
        except FileNotFoundError:
            logger.error(f"源目录不存在: {src_path}")
            raise
        except FileExistsError:
            logger.error(f"目标目录已存在: {dst_path}")
            raise
        except Exception as e:
            logger.error(f"复制目录失败: {src_path} -> {dst_path}, 错误: {str(e)}")
            raise

    @staticmethod
    def move_directory(src_path, dst_path):
        """
        移动目录
        
        Args:
            src_path: 源目录路径
            dst_path: 目标目录路径
        """
        # 直接使用已导入的config_manager
        # 处理相对路径
        if not os.path.isabs(src_path):
            src_path = os.path.join(config_manager.DATA_DIR, src_path)
        if not os.path.isabs(dst_path):
            dst_path = os.path.join(config_manager.DATA_DIR, dst_path)

        logger.debug(f"移动目录: {src_path} -> {dst_path}")
        try:
            # 确保目标目录的父目录存在
            os.makedirs(os.path.dirname(dst_path), exist_ok = True)
            shutil.move(src_path, dst_path)
            logger.info(f"目录移动成功: {src_path} -> {dst_path}")
        except FileNotFoundError:
            logger.error(f"源目录不存在: {src_path}")
            raise
        except Exception as e:
            logger.error(f"移动目录失败: {src_path} -> {dst_path}, 错误: {str(e)}")
            raise

    @staticmethod
    def read_text(file_path, encoding = 'utf-8'):
        """
        读取文本文件内容
        
        Args:
            file_path: 文件路径
            encoding: 文件编码,默认为utf-8
            
        Returns:
            str: 文件内容
        """
        # 直接使用已导入的config_manager
        # 处理相对路径
        if not os.path.isabs(file_path):
            file_path = os.path.join(config_manager.DATA_DIR, file_path)

        logger.debug(f"读取文本文件: {file_path}")
        try:
            with open(file_path, 'r', encoding = encoding) as f:
                content = f.read()
            logger.info(f"文本文件读取成功: {file_path}")
            return content
        except FileNotFoundError:
            logger.error(f"文件不存在: {file_path}")
            raise
        except UnicodeDecodeError:
            logger.error(f"文件编码错误: {file_path}, 尝试使用的编码: {encoding}")
            raise
        except Exception as e:
            logger.error(f"读取文本文件失败: {file_path}, 错误: {str(e)}")
            raise

    @staticmethod
    def write_text(file_path, content, encoding = 'utf-8'):
        """
        写入文本文件内容
        
        Args:
            file_path: 文件路径
            content: 要写入的内容
            encoding: 文件编码,默认为utf-8
        """
        # 直接使用已导入的config_manager
        # 处理相对路径
        if not os.path.isabs(file_path):
            file_path = os.path.join(config_manager.DATA_DIR, file_path)

        # 确保目录存在
        os.makedirs(os.path.dirname(file_path), exist_ok = True)

        logger.debug(f"写入文本文件: {file_path}")
        try:
            with open(file_path, 'w', encoding = encoding) as f:
                f.write(content)
            logger.info(f"文本文件写入成功: {file_path}")
        except Exception as e:
            logger.error(f"写入文本文件失败: {file_path}, 错误: {str(e)}")
            raise

    @staticmethod
    def append_text(file_path, content, encoding = 'utf-8', newline = '\n'):
        """
        向文本文件追加内容
        
        Args:
            file_path: 文件路径
            content: 要追加的内容
            encoding: 文件编码,默认为utf-8
            newline: 换行符,默认为'\n'
        """
        # 直接使用已导入的config_manager
        # 处理相对路径
        if not os.path.isabs(file_path):
            file_path = os.path.join(config_manager.DATA_DIR, file_path)

        # 确保目录存在
        os.makedirs(os.path.dirname(file_path), exist_ok = True)

        logger.debug(f"追加文本文件: {file_path}")
        try:
            # 检查文件是否存在,如存在且不为空则添加换行符
            if os.path.exists(file_path) and os.path.getsize(file_path) > 0:
                content = newline + content

            with open(file_path, 'a', encoding = encoding) as f:
                f.write(content)
            logger.info(f"文本文件追加成功: {file_path}")
        except Exception as e:
            logger.error(f"追加文本文件失败: {file_path}, 错误: {str(e)}")
            raise

    @staticmethod
    def read_lines(file_path, encoding = 'utf-8', strip = True):
        """
        按行读取文本文件
        
        Args:
            file_path: 文件路径
            encoding: 文件编码,默认为utf-8
            strip: 是否去除每行首尾的空白字符,默认为True
            
        Returns:
            list: 包含每行内容的列表
        """
        # 直接使用已导入的config_manager
        # 处理相对路径
        if not os.path.isabs(file_path):
            file_path = os.path.join(config_manager.DATA_DIR, file_path)

        logger.debug(f"按行读取文本文件: {file_path}")
        try:
            lines = []
            with open(file_path, 'r', encoding = encoding) as f:
                for line in f:
                    if strip:
                        line = line.strip()
                    lines.append(line)
            logger.info(f"文本文件按行读取成功: {file_path}")
            return lines
        except FileNotFoundError:
            logger.error(f"文件不存在: {file_path}")
            raise
        except UnicodeDecodeError:
            logger.error(f"文件编码错误: {file_path}, 尝试使用的编码: {encoding}")
            raise
        except Exception as e:
            logger.error(f"按行读取文本文件失败: {file_path}, 错误: {str(e)}")
            raise

    @staticmethod
    def get_file_size(file_path):
        """
        获取文件大小（字节数）
        
        Args:
            file_path: 文件路径
            
        Returns:
            int: 文件大小（字节）
        """
        # 直接使用已导入的config_manager
        if not os.path.isabs(file_path):
            file_path = os.path.join(config_manager.DATA_DIR, file_path)

        logger.debug(f"获取文件大小: {file_path}")
        try:
            size = os.path.getsize(file_path)
            logger.info(f"文件大小获取成功: {file_path}, 大小: {size} 字节")
            return size
        except FileNotFoundError:
            logger.error(f"文件不存在: {file_path}")
            raise
        except Exception as e:
            logger.error(f"获取文件大小失败: {file_path}, 错误: {str(e)}")
            raise

    @staticmethod
    def format_size(size_bytes):
        """
        格式化文件大小（将字节转换为KB/MB/GB等）
        
        Args:
            size_bytes: 文件大小（字节）
            
        Returns:
            str: 格式化后的文件大小
        """
        # 定义单位
        units = ['B', 'KB', 'MB', 'GB', 'TB']

        # 转换单位
        size = float(size_bytes)
        unit_index = 0
        while size >= 1024 and unit_index < len(units) - 1:
            size /= 1024
            unit_index += 1

        # 返回格式化的大小
        if unit_index == 0:
            return f"{int(size)} {units[unit_index]}"
        else:
            return f"{size:.2f} {units[unit_index]}"

    @staticmethod
    def get_file_info(file_path):
        """
        获取文件信息
        
        Args:
            file_path: 文件路径
            
        Returns:
            dict: 文件信息字典,包含大小、创建时间、修改时间等
        """
        # 直接使用已导入的config_manager
        # 处理相对路径
        if not os.path.isabs(file_path):
            file_path = os.path.join(config_manager.DATA_DIR, file_path)

        logger.debug(f"获取文件信息: {file_path}")
        try:
            # 获取文件状态
            stat_info = os.stat(file_path)

            # 构建文件信息字典
            file_info = {
                'path': file_path,
                'size': stat_info.st_size,
                'size_formatted': FileHandler.format_size(stat_info.st_size),
                'created_time': stat_info.st_ctime,
                'created_time_str': time.strftime('%Y-%m-%d %H:%M:%S', time.localtime(stat_info.st_ctime)),
                'modified_time': stat_info.st_mtime,
                'modified_time_str': time.strftime('%Y-%m-%d %H:%M:%S', time.localtime(stat_info.st_mtime)),
                'access_time': stat_info.st_atime,
                'access_time_str': time.strftime('%Y-%m-%d %H:%M:%S', time.localtime(stat_info.st_atime)),
                'is_file': os.path.isfile(file_path),
                'is_directory': os.path.isdir(file_path),
                'extension': os.path.splitext(file_path)[1].lower(),
                'name': os.path.basename(file_path)
            }

            logger.info(f"文件信息获取成功: {file_path}")
            return file_info
        except FileNotFoundError:
            logger.error(f"文件不存在: {file_path}")
            raise
        except Exception as e:
            logger.error(f"获取文件信息失败: {file_path}, 错误: {str(e)}")
            raise

    @staticmethod
    def is_file_exists(file_path):
        """
        检查文件或目录是否存在
        
        Args:
            file_path: 文件或目录路径
            
        Returns:
            bool: 存在返回True,否则返回False
        """
        # 直接使用已导入的config_manager
        # 处理相对路径
        if not os.path.isabs(file_path):
            file_path = os.path.join(config_manager.DATA_DIR, file_path)

        exists = os.path.exists(file_path)
        logger.debug(f"检查文件/目录是否存在: {file_path}, 结果: {exists}")
        return exists

    @staticmethod
    def get_file_extension(file_path):
        """
        获取文件扩展名
        
        Args:
            file_path: 文件路径
            
        Returns:
            str: 文件扩展名（包含点号,小写）
        """
        extension = os.path.splitext(file_path)[1].lower()
        logger.debug(f"获取文件扩展名: {file_path}, 扩展名: {extension}")
        return extension

    @staticmethod
    def calculate_file_hash(file_path, algorithm = 'md5', chunk_size = 8192):
        """
        计算文件哈希值
        
        Args:
            file_path: 文件路径
            algorithm: 哈希算法,可选值: 'md5', 'sha1', 'sha256'
            chunk_size: 读取文件的块大小
            
        Returns:
            str: 文件哈希值（十六进制字符串）
        """
        # 直接使用已导入的config_manager
        # 处理相对路径
        if not os.path.isabs(file_path):
            file_path = os.path.join(config_manager.DATA_DIR, file_path)

        logger.debug(f"计算文件哈希值: {file_path}, 算法: {algorithm}")
        try:
            # 选择哈希算法
            if algorithm.lower() == 'md5':
                hash_obj = hashlib.md5()
            elif algorithm.lower() == 'sha1':
                hash_obj = hashlib.sha1()
            elif algorithm.lower() == 'sha256':
                hash_obj = hashlib.sha256()
            else:
                raise ValueError(f"不支持的哈希算法: {algorithm}")

            # 分块读取文件并计算哈希值
            with open(file_path, 'rb') as f:
                while chunk := f.read(chunk_size):
                    hash_obj.update(chunk)

            # 获取十六进制哈希值
            hash_value = hash_obj.hexdigest()
            logger.info(f"文件哈希值计算成功: {file_path}, {algorithm.upper()}: {hash_value}")
            return hash_value
        except FileNotFoundError:
            logger.error(f"文件不存在: {file_path}")
            raise
        except Exception as e:
            logger.error(f"计算文件哈希值失败: {file_path}, 算法: {algorithm}, 错误: {str(e)}")
            raise

    @staticmethod
    def calculate_file_md5(file_path, chunk_size = 8192):
        """
        计算文件MD5哈希值
        
        Args:
            file_path: 文件路径
            chunk_size: 读取文件的块大小
            
        Returns:
            str: MD5哈希值
        """
        return FileHandler.calculate_file_hash(file_path, algorithm = 'md5', chunk_size = chunk_size)

    @staticmethod
    def calculate_file_sha1(file_path, chunk_size = 8192):
        """
        计算文件SHA1哈希值
        
        Args:
            file_path: 文件路径
            chunk_size: 读取文件的块大小
            
        Returns:
            str: SHA1哈希值
        """
        return FileHandler.calculate_file_hash(file_path, algorithm = 'sha1', chunk_size = chunk_size)

    @staticmethod
    def calculate_file_sha256(file_path, chunk_size = 8192):
        """
        计算文件SHA256哈希值
        
        Args:
            file_path: 文件路径
            chunk_size: 读取文件的块大小
            
        Returns:
            str: SHA256哈希值
        """
        return FileHandler.calculate_file_hash(file_path, algorithm = 'sha256', chunk_size = chunk_size)

    @staticmethod
    def compare_file_hash(file_path, expected_hash, algorithm = 'md5'):
        """
        比较文件哈希值
        
        Args:
            file_path: 文件路径
            expected_hash: 预期的哈希值
            algorithm: 哈希算法
            
        Returns:
            bool: 哈希值匹配返回True,否则返回False
        """
        actual_hash = FileHandler.calculate_file_hash(file_path, algorithm = algorithm)
        result = actual_hash.lower() == expected_hash.lower()
        logger.debug(f"文件哈希值比较结果: {file_path}, 匹配: {result}")
        return result

    @staticmethod
    def read_binary(file_path):
        """
        读取二进制文件
        
        Args:
            file_path: 文件路径
            
        Returns:
            bytes: 二进制数据
        """
        # 直接使用已导入的config_manager
        # 处理相对路径
        if not os.path.isabs(file_path):
            file_path = os.path.join(config_manager.DATA_DIR, file_path)

        logger.debug(f"读取二进制文件: {file_path}")
        try:
            with open(file_path, 'rb') as f:
                data = f.read()
            logger.info(f"二进制文件读取成功: {file_path}, 数据大小: {len(data)}字节")
            return data
        except FileNotFoundError:
            logger.error(f"文件不存在: {file_path}")
            raise
        except Exception as e:
            logger.error(f"读取二进制文件失败: {file_path}, 错误: {str(e)}")
            raise

    @staticmethod
    def write_binary(file_path, data):
        """
        写入二进制文件
        
        Args:
            file_path: 文件路径
            data: 二进制数据（bytes类型）
        """
        # 直接使用已导入的config_manager
        # 处理相对路径
        if not os.path.isabs(file_path):
            file_path = os.path.join(config_manager.DATA_DIR, file_path)

        # 确保目录存在
        os.makedirs(os.path.dirname(file_path), exist_ok = True)

        logger.debug(f"写入二进制文件: {file_path}, 数据大小: {len(data)}字节")
        try:
            with open(file_path, 'wb') as f:
                f.write(data)
            logger.info(f"二进制文件写入成功: {file_path}, 数据大小: {len(data)}字节")
        except Exception as e:
            logger.error(f"写入二进制文件失败: {file_path}, 错误: {str(e)}")
            raise

    @staticmethod
    def append_binary(file_path, data):
        """
        追加二进制数据到文件
        
        Args:
            file_path: 文件路径
            data: 二进制数据（bytes类型）
        """
        # 直接使用已导入的config_manager
        # 处理相对路径
        if not os.path.isabs(file_path):
            file_path = os.path.join(config_manager.DATA_DIR, file_path)

        # 确保目录存在
        os.makedirs(os.path.dirname(file_path), exist_ok = True)

        logger.debug(f"追加二进制数据到文件: {file_path}, 追加数据大小: {len(data)}字节")
        try:
            with open(file_path, 'ab') as f:
                f.write(data)
            logger.info(f"二进制数据追加成功: {file_path}, 追加数据大小: {len(data)}字节")
        except Exception as e:
            logger.error(f"追加二进制数据失败: {file_path}, 错误: {str(e)}")
            raise

    @staticmethod
    def read_binary_chunked(file_path, chunk_size = 8192, callback = None):
        """
        分块读取二进制文件
        
        Args:
            file_path: 文件路径
            chunk_size: 块大小
            callback: 可选的回调函数,接收块数据和块索引
        """
        # 直接使用已导入的config_manager
        # 处理相对路径
        if not os.path.isabs(file_path):
            file_path = os.path.join(config_manager.DATA_DIR, file_path)

        logger.debug(f"分块读取二进制文件: {file_path}, 块大小: {chunk_size}")
        try:
            with open(file_path, 'rb') as f:
                chunk_index = 0
                while True:
                    chunk = f.read(chunk_size)
                    if not chunk:
                        break
                    if callback:
                        callback(chunk, chunk_index)
                    chunk_index += 1
            logger.info(f"分块读取二进制文件完成: {file_path}, 总块数: {chunk_index}")
        except FileNotFoundError:
            logger.error(f"文件不存在: {file_path}")
            raise
        except Exception as e:
            logger.error(f"分块读取二进制文件失败: {file_path}, 错误: {str(e)}")
            raise

    @staticmethod
    def write_binary_chunked(file_path, chunks, chunk_size = 8192):
        """
        分块写入二进制文件
        
        Args:
            file_path: 文件路径
            chunks: 包含二进制数据的可迭代对象
            chunk_size: 块大小（用于日志记录）
        """
        # 直接使用已导入的config_manager
        # 处理相对路径
        if not os.path.isabs(file_path):
            file_path = os.path.join(config_manager.DATA_DIR, file_path)

        # 确保目录存在
        os.makedirs(os.path.dirname(file_path), exist_ok = True)

        logger.debug(f"分块写入二进制文件: {file_path}, 块大小: {chunk_size}")
        try:
            with open(file_path, 'wb') as f:
                for chunk_index, chunk in enumerate(chunks):
                    f.write(chunk)
            logger.info(f"分块写入二进制文件完成: {file_path}, 总块数: {chunk_index + 1}")
        except Exception as e:
            logger.error(f"分块写入二进制文件失败: {file_path}, 错误: {str(e)}")
            raise


# 定义DataHandler作为FileHandler的别名,保持代码兼容性
DataHandler = FileHandler
